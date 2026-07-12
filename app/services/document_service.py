"""
Document parsing service.
Supports: .pdf (PyMuPDF), .docx (python-docx), .txt (plain text)
"""

import html as html_lib
import io
from pathlib import Path
from typing import List, Optional

import fitz  # PyMuPDF
import docx


class DocumentService:

    def parse_file(self, content: bytes, filename: str) -> str:
        return self.parse_file_with_pages(content, filename)["text"]

    def parse_file_with_pages(self, content: bytes, filename: str) -> dict:
        """Parse file and return text + per-page breakdown for citation mapping."""
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return self._parse_pdf_with_pages(content, filename)
        elif ext in (".docx", ".doc"):
            text = self._parse_docx(content)
            return {"text": text, "pages": [{"page": 1, "text": text}], "type": "docx", "filename": filename}
        elif ext == ".txt":
            text = content.decode("utf-8", errors="replace")
            return {"text": text, "pages": [{"page": 1, "text": text}], "type": "txt", "filename": filename}
        else:
            raise ValueError(f"Unsupported file type: {ext}. Use PDF, DOCX, or TXT.")

    def _parse_pdf_with_pages(self, content: bytes, filename: str) -> dict:
        doc = fitz.open(stream=content, filetype="pdf")
        pages = []
        all_text = []
        for i, page in enumerate(doc):
            text = page.get_text()
            pages.append({"page": i + 1, "text": text})
            all_text.append(text)
        return {"text": "\n\n".join(all_text), "pages": pages, "type": "pdf", "filename": filename}

    def _parse_docx(self, content: bytes) -> str:
        doc = docx.Document(io.BytesIO(content))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

    def find_quote_page(self, pages: List[dict], quote: str) -> Optional[int]:
        """Return 1-based page number where quote appears, or None."""
        if not quote or not pages:
            return None
        q = quote.lower().strip()
        for page_info in pages:
            if q in page_info["text"].lower():
                return page_info["page"]
        # Partial match — first 60 non-trivial characters
        short = q[:60].strip()
        if len(short) >= 15:
            for page_info in pages:
                if short in page_info["text"].lower():
                    return page_info["page"]
        return None

    def to_html(self, content: bytes, filename: str) -> str:
        """Convert document to HTML string for in-browser preview."""
        ext = Path(filename).suffix.lower()
        if ext in (".docx", ".doc"):
            return self._docx_to_html(content)
        elif ext == ".txt":
            return self._txt_to_html(content.decode("utf-8", errors="replace"))
        return "<html><body><p>Preview not available for this file type.</p></body></html>"

    def _docx_to_html(self, content: bytes) -> str:
        doc = docx.Document(io.BytesIO(content))
        styles = (
            "<style>"
            "body{font-family:Georgia,serif;padding:2rem 2.5rem;max-width:760px;"
            "margin:0 auto;color:#1a1a1a;line-height:1.75;font-size:15px;background:#fff;}"
            "h1{font-size:1.55rem;font-weight:700;margin:1.75rem 0 0.75rem;}"
            "h2{font-size:1.25rem;font-weight:600;margin:1.5rem 0 0.6rem;}"
            "h3{font-size:1.05rem;font-weight:600;margin:1.25rem 0 0.5rem;}"
            "p{margin:0.45rem 0;}"
            "mark{background:#fff59d;border-radius:2px;padding:0 3px;}"
            "</style>"
        )
        parts = [f"<html><head><meta charset='utf-8'>{styles}</head><body>"]
        for para in doc.paragraphs:
            if not para.text.strip():
                parts.append("<br>")
                continue
            style_name = (para.style.name or "").lower()
            escaped = html_lib.escape(para.text)
            if "heading 1" in style_name:
                parts.append(f"<h1>{escaped}</h1>")
            elif "heading 2" in style_name:
                parts.append(f"<h2>{escaped}</h2>")
            elif "heading 3" in style_name or "heading 4" in style_name:
                parts.append(f"<h3>{escaped}</h3>")
            else:
                parts.append(f"<p>{escaped}</p>")
        parts.append("</body></html>")
        return "\n".join(parts)

    def _txt_to_html(self, text: str) -> str:
        escaped = html_lib.escape(text)
        return (
            "<html><head><meta charset='utf-8'><style>"
            "body{font-family:'Courier New',monospace;padding:2rem;white-space:pre-wrap;"
            "color:#1a1a1a;line-height:1.65;font-size:13px;background:#fff;}"
            "mark{background:#fff59d;border-radius:2px;padding:0 2px;}"
            f"</style></head><body>{escaped}</body></html>"
        )
